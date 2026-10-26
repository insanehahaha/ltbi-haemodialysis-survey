# 04_tablolar.py
# python 3.11.17, paket sürümleri requirements.txt dosyasında
#
# bu dosya makaledeki dört tabloyu word dosyası olarak hazırlar
# tablolar ingilizce, önceki makalelerimizdeki biçimde:
#   Times New Roman, gövde 8 punto
#   sadece üç yatay çizgi (en üst, başlık satırının altı, en alt), dikey çizgi yok
#   başlık satırı kalın, bölüm satırları bütün sütunları kaplayan birleşik satır
#   tablo adı tablonun üstünde, açıklamalar ve kısaltmalar altında
#
# sayılar kodlanmış veriden ve 02 ile 03'ün kaydettiği sonuçlardan okunuyor
#
# çıktı: çıktılar/tablolar/Tablolar.docx (her tablo ayrı sayfada)

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

kod_klasoru = Path(__file__).resolve().parent
veri_klasoru = kod_klasoru.parent / "veri"
tablo_klasoru = kod_klasoru.parent / "tablolar"
tablo_klasoru.mkdir(exist_ok=True)

veri = pd.read_csv(veri_klasoru / "kodlanmis_veri.csv")
s02 = json.load(open(veri_klasoru / "sayilar_02.json", encoding="utf-8"))
s03 = json.load(open(veri_klasoru / "sayilar_03.json", encoding="utf-8"))
maddeler = pd.read_excel(veri_klasoru / "tanimlayici_sonuclar.xlsx", sheet_name="T2_maddeler").set_index("madde")
modeller = pd.read_excel(veri_klasoru / "model_sonuclari.xlsx", sheet_name=None, index_col=0)

FONT = "Times New Roman"
brans_sira = ["Nefroloji", "İç hastalıkları", "Pratisyen hekim"]


# ---------------------------------------------------------------
# sayı biçimlendirme
# ---------------------------------------------------------------
# yüzdeler bir ondalık, p değerleri üç ondalık, 0.001'in altı "<0.001"
# ondalık ayıracı nokta

def yuzde_yaz(n, payda):
    if payda == 0:
        return "–"
    return f"{n} ({n / payda * 100:.1f})"


def p_yaz(p):
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def eksi(metin):
    # tire yerine eksi işareti
    return metin.replace("-", "−")


# ---------------------------------------------------------------
# word tablosu için yardımcı fonksiyonlar
# ---------------------------------------------------------------

def hucre_kenari(hucre, ust=None, alt=None):
    # hücrenin üst ve alt çizgisini ayarlıyor
    tcPr = hucre._tc.get_or_add_tcPr()
    kenarlar = tcPr.find(qn("w:tcBorders"))
    if kenarlar is None:
        kenarlar = OxmlElement("w:tcBorders")
        tcPr.append(kenarlar)
    for yon, kalinlik in [("top", ust), ("bottom", alt)]:
        if kalinlik is None:
            continue
        cizgi = OxmlElement(f"w:{yon}")
        cizgi.set(qn("w:val"), "single")
        cizgi.set(qn("w:sz"), str(kalinlik))
        cizgi.set(qn("w:space"), "0")
        cizgi.set(qn("w:color"), "000000")
        kenarlar.append(cizgi)


def tablo_kenarlarini_kaldir(tablo):
    # varsayılan bütün çizgileri kaldırıyor
    tblPr = tablo._tbl.tblPr
    kenarlar = OxmlElement("w:tblBorders")
    for yon in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        cizgi = OxmlElement(f"w:{yon}")
        cizgi.set(qn("w:val"), "nil")
        kenarlar.append(cizgi)
    tblPr.append(kenarlar)


def yazi_yaz(hucre, metin, kalin=False, italik=False, hizalama="sol", punto=8):
    hucre.text = ""
    paragraf = hucre.paragraphs[0]
    paragraf.paragraph_format.space_before = Pt(1)
    paragraf.paragraph_format.space_after = Pt(1)
    paragraf.alignment = WD_ALIGN_PARAGRAPH.LEFT if hizalama == "sol" else WD_ALIGN_PARAGRAPH.CENTER
    parca = paragraf.add_run(metin)
    parca.font.name = FONT
    parca._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    parca.font.size = Pt(punto)
    parca.bold = kalin
    parca.italic = italik


def paragraf_yaz(belge, parcalar, punto=10, sonra=4):
    # parcalar: [(metin, kalin_mi), ...]
    paragraf = belge.add_paragraph()
    paragraf.paragraph_format.space_after = Pt(sonra)
    for metin, kalin in parcalar:
        parca = paragraf.add_run(metin)
        parca.font.name = FONT
        parca.font.size = Pt(punto)
        parca.bold = kalin
    return paragraf


def tablo_olustur(belge, numara, baslik, sutun_basliklari, satirlar, dipnot, sutun_genislikleri):
    # satır türleri:
    #   ("bolum", "Bölüm adı")       -> bütün sütunları kaplayan birleşik satır
    #   ("veri", [hücre1, hücre2..]) -> normal satır, ilk hücre sola diğerleri ortaya hizalı
    #   ("veri_girinti", [...])      -> ilk hücresi içeriden başlayan normal satır
    #   ("alt_baslik", [...])        -> tablonun ortasında kalın yazılan yeni sütun başlıkları
    paragraf_yaz(belge, [(f"Table {numara}. ", True), (baslik, False)])

    sutun_sayisi = len(sutun_basliklari)
    tablo = belge.add_table(rows=1, cols=sutun_sayisi)
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    tablo.autofit = False
    tablo_kenarlarini_kaldir(tablo)

    for i, metin in enumerate(sutun_basliklari):
        hucre = tablo.rows[0].cells[i]
        yazi_yaz(hucre, metin, kalin=True, hizalama="sol" if i == 0 else "orta")
        hucre_kenari(hucre, ust=8, alt=4)

    for tur, icerik in satirlar:
        yeni = tablo.add_row().cells
        if tur == "bolum":
            birlesik = yeni[0].merge(yeni[-1])
            yazi_yaz(birlesik, icerik, kalin=True)
        else:
            for i, metin in enumerate(icerik):
                if i == 0 and tur == "veri_girinti":
                    metin = "   " + metin
                yazi_yaz(yeni[i], metin, kalin=(tur == "alt_baslik"), hizalama="sol" if i == 0 else "orta")

    for hucre in tablo.rows[-1].cells:
        hucre_kenari(hucre, alt=8)

    # sütun genişliklerini sabitliyoruz
    duzen = OxmlElement("w:tblLayout")
    duzen.set(qn("w:type"), "fixed")
    tablo._tbl.tblPr.append(duzen)
    for i, genislik in enumerate(sutun_genislikleri):
        tablo.columns[i].width = Cm(genislik)
    for satir in tablo.rows:
        for i, genislik in enumerate(sutun_genislikleri):
            satir.cells[i].width = Cm(genislik)

    # tablo sonraki sayfaya taşarsa başlık satırı tekrar ediyor
    trPr = tablo.rows[0]._tr.get_or_add_trPr()
    tekrar = OxmlElement("w:tblHeader")
    tekrar.set(qn("w:val"), "true")
    trPr.append(tekrar)

    paragraf = paragraf_yaz(belge, [(dipnot, False)], punto=8, sonra=0)
    paragraf.paragraph_format.space_before = Pt(3)
    return tablo


# her tablo yeni bir bölümde, yatay ya da dikey sayfada başlıyor
def yeni_bolum(belge, yatay=False):
    bolum = belge.add_section(WD_SECTION.NEW_PAGE)
    if yatay:
        bolum.orientation = WD_ORIENT.LANDSCAPE
        bolum.page_width, bolum.page_height = Cm(29.7), Cm(21.0)
    else:
        bolum.orientation = WD_ORIENT.PORTRAIT
        bolum.page_width, bolum.page_height = Cm(21.0), Cm(29.7)


belge = Document()
bolum = belge.sections[0]
bolum.orientation = WD_ORIENT.PORTRAIT
bolum.page_width, bolum.page_height = Cm(21.0), Cm(29.7)
for kenar in ["left_margin", "right_margin", "top_margin", "bottom_margin"]:
    setattr(bolum, kenar, Cm(2.0))


# ---------------------------------------------------------------
# Table 1: katılımcı ve ünite özellikleri, branşa göre
# ---------------------------------------------------------------
# her değişken için: genel ve branşlara göre n (%)
# yüzdenin paydası o soruyu yanıtlayanlar

brans_n = {b: int((veri["brans"] == b).sum()) for b in brans_sira}

tablo1_icerik = [
    ("Years in medical practice", "hekimlik_yili",
     [("0–5", "0-5"), ("6–10", "6-10"), ("11–20", "11-20"), (">20", ">20")]),
    ("Years of haemodialysis experience", "hd_deneyimi",
     [("<1", "<1"), ("1–3", "1-3"), ("4–10", "4-10"), (">10", ">10")]),
    ("Type of centre", "merkez_4grup",
     [("University hospital", "Üniversite"), ("Training and research hospital", "Eğitim-Araştırma"),
      ("State hospital", "Devlet"), ("Private hospital or dialysis centre", "Özel")]),
    ("Approximate number of patients in the unit", "hasta_sayisi",
     [("<50", "<50"), ("50–100", "50-100"), ("101–200", "101-200"), (">200", ">200")]),
    ("Separate area or room for a patient with suspected TB", "izolasyon_alani",
     [("Yes", "Evet"), ("No", "Hayır"), ("Don't know", "Bilmiyorum")]),
    ("TB/LTBI training in the past 2 years", "egitim",
     [("Yes", "Evet"), ("No", "Hayır")]),
    ("Patient diagnosed with active TB in the unit, past 12 months", "tb_olgusu",
     [("Yes", "Evet"), ("No", "Hayır"), ("Don't know", "Bilmiyorum")]),
    ("Patient masked or isolated for suspected active TB, past 12 months", "tb_izolasyon",
     [("Yes", "Evet"), ("No", "Hayır"), ("Don't know", "Bilmiyorum")]),
    ("Transplant candidates evaluated in the unit", "transplant_degerlendirme",
     [("Yes, routinely", "Evet, rutin"), ("Yes, sometimes", "Evet, bazen"), ("No", "Hayır"),
      ("Don't know", "Bilmiyorum")]),
    ("Who most often initiates LTBI screening in the unit", "tarama_baslatan",
     [("Nephrologist/HD physician", "Nefroloji/HD hekimi"), ("Pulmonology", "Göğüs hastalıkları"),
      ("Infectious diseases", "Enfeksiyon hastalıkları"), ("More than one", "Birden fazla aktör"),
      ("Institutional protocol (automatic)", "Kurum protokolü (otomatik)"), ("Varies", "Değişken"),
      ("Don't know", "Bilmiyorum")]),
    ("Written TB/LTBI protocol in the unit", "yazili_protokol",
     [("Yes", "Evet"), ("No", "Hayır"), ("Don't know", "Bilmiyorum")]),
]

satirlar = []
for bolum_adi, sutun, kategoriler in tablo1_icerik:
    satirlar.append(("bolum", bolum_adi))
    genel_payda = veri[sutun].notna().sum()
    for etiket, kod in kategoriler:
        hucreler = [etiket, yuzde_yaz(int((veri[sutun] == kod).sum()), genel_payda)]
        for b in brans_sira:
            alt = veri[veri["brans"] == b][sutun]
            hucreler.append(yuzde_yaz(int((alt == kod).sum()), alt.notna().sum()))
        satirlar.append(("veri_girinti", hucreler))

tablo_olustur(
    belge, 1,
    "Characteristics of the participants and their haemodialysis units, overall and by specialty",
    ["Characteristic",
     f"All\n(n = {len(veri)})",
     f"Nephrologists\n(n = {brans_n['Nefroloji']})",
     f"Internists\n(n = {brans_n['İç hastalıkları']})",
     f"General practitioners\n(n = {brans_n['Pratisyen hekim']})"],
    satirlar,
    "Values are n (%). Percentages are calculated among participants who answered each question. "
    "Private hospital and private dialysis centre were combined. "
    "HD, haemodialysis; LTBI, latent tuberculosis infection; TB, tuberculosis.",
    [6.6, 2.6, 2.6, 2.6, 2.6],
)
yeni_bolum(belge, yatay=True)


# ---------------------------------------------------------------
# Table 2: bilgi ve senaryo maddeleri, branşa göre doğru yanıt
# bu tablo yatay sayfada
# ---------------------------------------------------------------
madde_metinleri = {
    "D1": ("Test(s) used for LTBI screening", "TST and IGRA"),
    "D2": ("Most accurate statement on IGRA and TST",
           "Both support LTBI; active TB requires clinical, radiological and microbiological evaluation"),
    "D3": ("HD/ESKD is a condition that increases the risk of progression from LTBI to active TB", "True"),
    "D4": ("HD patient group with the highest priority for LTBI screening",
           "Transplant candidates or planned immunosuppression"),
    "D5": ("Approach generally preferred in guidelines for LTBI treatment",
           "Short rifamycin-based regimens; isoniazid if these are unsuitable"),
    "D6": ("Most important factor in the choice of LTBI treatment",
           "Drug interactions of rifamycins"),
    "D7": ("In an HD patient with suspected active TB, an LTBI test alone is sufficient to guide management",
           "False"),
    "D8": ("Starting LTBI treatment before active TB is excluded",
           "Active TB evaluation is required because of misdiagnosis and resistance risks"),
    "D9": ("First infection-control step for an HD patient with suspected active TB",
           "Mask or isolation and rapid referral according to institutional protocol"),
    "D10": ("Routine repetition of IGRA/TST is generally not recommended, but re-evaluation may be needed "
            "after a new exposure", "True"),
    "E1": ("Scenario 1: 58-year-old man on HD for 4 years with weight loss, night sweats and suspicious "
           "chest radiograph", "Start evaluation and isolation for active TB and refer"),
    "E2": ("Scenario 2: 45-year-old asymptomatic transplant candidate on HD, IGRA positive, no evidence of "
           "active TB", "Plan evaluation for LTBI treatment with infectious diseases or pulmonology"),
}


def madde_satiri(soru):
    m = maddeler.loc[soru]
    metin, anahtar = madde_metinleri[soru]
    hucreler = [f"{soru}. {metin}", anahtar,
                yuzde_yaz(int(m["dogru_n"]), int(m["yanitlayan"])),
                yuzde_yaz(int(m["bilmiyorum_n"]), int(m["yanitlayan"]))]
    for b in brans_sira:
        alt = veri[veri["brans"] == b][soru].dropna()
        hucreler.append(yuzde_yaz(int(alt.sum()), len(alt)))
    p_metni = p_yaz(m["brans_p"])
    if m["brans_test"] == "Fisher":
        p_metni += "ᵃ"
    hucreler.append(p_metni)
    return hucreler


def skor_satiri(etiket, skor, p):
    hucreler = [etiket, "", f"{veri[skor].mean():.2f} ± {veri[skor].std():.2f}", ""]
    for b in brans_sira:
        alt = veri[veri["brans"] == b][skor].dropna()
        hucreler.append(f"{alt.mean():.2f} ± {alt.std():.2f}")
    hucreler.append(p_yaz(p) + "ᵇ")
    return hucreler


satirlar = [("bolum", "Knowledge items")]
for soru in ["D1", "D2", "D3", "D4", "D5", "D6", "D7", "D8", "D9", "D10"]:
    satirlar.append(("veri", madde_satiri(soru)))
satirlar.append(("veri", skor_satiri("Knowledge score (0–10), mean ± SD", "bilgi_skoru",
                                     s02["bilgi|brans|kruskal_p"])))
satirlar.append(("bolum", "Clinical scenarios"))
for soru in ["E1", "E2"]:
    satirlar.append(("veri", madde_satiri(soru)))
satirlar.append(("veri", skor_satiri("Scenario score (0–2), mean ± SD", "senaryo_skoru",
                                     s02["aile|brans|senaryo_skoru|p"])))

tablo_olustur(
    belge, 2,
    "Correct answers to the knowledge items and clinical scenarios, overall and by specialty",
    ["Item", "Correct answer", "Correct,\nall", "\"Don't know\",\nall",
     "Nephrologists,\ncorrect", "Internists,\ncorrect", "General practitioners,\ncorrect", "p"],
    satirlar,
    "Values are n (%) unless otherwise stated. Percentages are calculated among participants who answered "
    "each item. An item was scored as correct when the correct option was marked alone or together with "
    "one other option; for D1, marking TST and IGRA together was scored as correct. "
    f"Internal consistency of the knowledge score: KR-20 = {s02['kr20']:.2f}. "
    "p values compare the proportion of correct answers across specialties with the χ² test, "
    "or ᵃ with the Fisher–Freeman–Halton exact test when any expected count was below 5; "
    "ᵇ Kruskal–Wallis test. "
    "ESKD, end-stage kidney disease; HD, haemodialysis; IGRA, interferon-gamma release assay; "
    "KR-20, Kuder–Richardson formula 20; LTBI, latent tuberculosis infection; SD, standard deviation; "
    "TB, tuberculosis; TST, tuberculin skin test.",
    [7.0, 5.6, 2.0, 2.0, 2.4, 2.0, 2.4, 1.4],
)
yeni_bolum(belge)


# ---------------------------------------------------------------
# Table 3: çok değişkenli modeller
# ---------------------------------------------------------------
# bilgi skoru için üç doğrusal model ve tarama için Firth modeli

def ga_yaz(alt, ust, ondalik=2):
    return eksi(f"{alt:.{ondalik}f} to {ust:.{ondalik}f}")


def ols_satiri(model_adi, katsayi, etiket):
    m = modeller[model_adi].loc[katsayi]
    return ("veri_girinti", [etiket, eksi(f"{m['beta']:.2f}"), ga_yaz(m["ga_alt"], m["ga_ust"]), p_yaz(m["p"])])


ortak_satirlar = [
    ("HD experience ≤3 years (ref. >10 years)", "≤3"),
    ("HD experience 4–10 years (ref. >10 years)", "4-10"),
    ("TB/LTBI training in the past 2 years", "egitim_var"),
    ("Written TB/LTBI protocol in the unit", "protokol_var"),
    ("Routine evaluation of transplant candidates in the unit", "transplant_rutin"),
]

satirlar = []
satirlar.append(("bolum", f"Knowledge score, linear regression — Model A: specialty "
                          f"(n = {s03['model_A|n']}; R² = {s03['model_A|r2']:.3f})"))
satirlar.append(ols_satiri("model_A", "İç hastalıkları", "Internist (ref. nephrologist)"))
satirlar.append(ols_satiri("model_A", "Pratisyen hekim", "General practitioner (ref. nephrologist)"))
for etiket, katsayi in ortak_satirlar:
    satirlar.append(ols_satiri("model_A", katsayi, etiket))

satirlar.append(("bolum", f"Knowledge score, linear regression — Model B: type of centre "
                          f"(n = {s03['model_B|n']}; R² = {s03['model_B|r2']:.3f})"))
satirlar.append(ols_satiri("model_B", "Devlet", "State hospital (ref. academic)"))
satirlar.append(ols_satiri("model_B", "Özel", "Private (ref. academic)"))
for etiket, katsayi in ortak_satirlar:
    satirlar.append(ols_satiri("model_B", katsayi, etiket))

satirlar.append(("bolum", f"Knowledge score, linear regression — Model C: working environment, exploratory "
                          f"(n = {s03['model_C|n']}; R² = {s03['model_C|r2']:.3f})"))
satirlar.append(ols_satiri("model_C", "Nephrologist, state/private",
                           "Nephrologist, state or private (ref. nephrologist, academic)"))
satirlar.append(ols_satiri("model_C", "Internist", "Internist (ref. nephrologist, academic)"))
satirlar.append(ols_satiri("model_C", "General practitioner",
                           "General practitioner (ref. nephrologist, academic)"))
for etiket, katsayi in ortak_satirlar:
    satirlar.append(ols_satiri("model_C", katsayi, etiket))

satirlar.append(("bolum", "Requested LTBI screening in the past 12 months, Firth logistic regression "
                          f"(n = {s03['firth|n']}; {s03['firth|olay']} events)"))
satirlar.append(("alt_baslik", ["", "OR", "95% CI", "p"]))
firth_etiketleri = [
    ("Active TB case in the unit, past 12 months", "tb_olgusu_var"),
    ("Nephrologist (ref. internist or general practitioner)", "nefrolog"),
    ("Knowledge score, per 1 point", "bilgi_skoru"),
    ("Barrier score, per 1 point", "bariyer"),
]
for etiket, degisken in firth_etiketleri:
    m = modeller["firth"].loc[degisken]
    satirlar.append(("veri_girinti", [etiket, f"{m['OR']:.2f}", ga_yaz(m["ga_alt"], m["ga_ust"]), p_yaz(m["p"])]))

tablo_olustur(
    belge, 3,
    "Factors associated with the knowledge score and with having requested LTBI screening",
    ["Variable", "β", "95% CI", "p"],
    satirlar,
    "β is the difference in knowledge score (0–10 points) compared with the reference category. "
    "Models A, B and C were fitted separately because specialty and type of centre were strongly "
    f"associated (Cramér's V = {s02['cramer_v_brans_merkez']:.2f}); all three were adjusted for the same "
    "covariates and fitted in the same participants with complete data. Linear models used "
    "heteroscedasticity-consistent (HC3) standard errors. Academic centres include university and training "
    "and research hospitals. The outcome of the Firth penalized logistic regression model was having requested "
    "LTBI screening for at least one patient in the past 12 months; participants who did not remember the "
    "number of screenings were excluded. Confidence intervals are profile penalized likelihood intervals and "
    "p values are from penalized likelihood ratio tests. The barrier score ranges from 1 to 5, with higher "
    "scores indicating more barriers. "
    "CI, confidence interval; HD, haemodialysis; LTBI, latent tuberculosis infection; OR, odds ratio; "
    "ref., reference category; TB, tuberculosis.",
    [9.0, 2.0, 3.3, 1.7],
)
yeni_bolum(belge)


# ---------------------------------------------------------------
# Table 4: uygulama ve eğitim gereksinimi
# ---------------------------------------------------------------

def tek_soru_satirlari(sutun, kategoriler):
    payda = veri[sutun].notna().sum()
    return [("veri_girinti", [etiket, yuzde_yaz(int((veri[sutun] == kod).sum()), payda)])
            for etiket, kod in kategoriler]


def coklu_soru_satirlari(kukla_sutunlari):
    sonuc = []
    for etiket, sutun in kukla_sutunlari:
        payda = veri[sutun].notna().sum()
        sonuc.append(("veri_girinti", [etiket, yuzde_yaz(int(veri[sutun].sum()), payda)]))
    return sonuc


satirlar = []
satirlar.append(("bolum", "Number of HD patients for whom the physician initiated or requested LTBI screening, past 12 months"))
satirlar += tek_soru_satirlari("tarama_sayisi", [("0", "0"), ("1–5", "1-5"), ("6–19", "6-19"),
                                                  ("20–50", "20-50"), (">50", ">50"),
                                                  ("Do not remember", "Hatırlamıyorum")])
satirlar.append(("bolum", "Situations in which LTBI screening would be requested (multiple answers allowed)"))
satirlar += coklu_soru_satirlari([("Transplant candidate or planned immunosuppression", "F2_transplant"),
                                  ("History of TB contact", "F2_temas"),
                                  ("Active TB case in the unit", "F2_unitede_olgu"),
                                  ("Origin from or long stay in a high-incidence area", "F2_insidans"),
                                  ("Previous TB or TB sequelae", "F2_eski_tb"),
                                  ("Routinely in all patients", "F2_rutin")])
satirlar.append(("bolum", "Approach to a positive LTBI test"))
satirlar += tek_soru_satirlari("pozitiflikte_yaklasim",
                               [("Refer to infectious diseases or pulmonology",
                                 "Enfeksiyon hastalıkları /Göğüs Hastalıklarına yönlendiririm"),
                                ("Start LTBI treatment myself", "Kendim LTBI tedavisi başlarım"),
                                ("Follow up only", "Sadece izlerim"),
                                ("More than one answer", "Çoklu/değişken"),
                                ("Never encountered", "Hiç karşılaşmadım")])
satirlar.append(("bolum", "Having to manage a patient with suspected active TB"))
satirlar += tek_soru_satirlari("suphe_yonetimi",
                               [("Often", "Sık"), ("Sometimes", "Bazen"), ("Rarely", "Nadiren"),
                                ("Never", "Hiç"),
                                ("No suspected case in the past 12 months",
                                 "Son 12 ayda aktif TB şüpheli olgu ile karşılaşmadım")])
satirlar.append(("bolum", "Contact evaluation after suspected TB or exposure in the unit"))
satirlar += tek_soru_satirlari("temasli_degerlendirme",
                               [("Institutional protocol in place and applied", "Kurum protokolü var ve uygulanıyor"),
                                ("Depends on the responsible person", "Sorumlu kişiye göre değişiyor"),
                                ("Not performed", "Uygulanmıyor"),
                                ("More than one answer", "Çoklu/değişken"),
                                ("Not encountered in the past 12 months", "Son 12 ayda bu durumla karşılaşmadım"),
                                ("Don't know", "Bilmiyorum")])
satirlar.append(("bolum", "Clinical presentation that prompts faster evaluation for active TB"))
satirlar += tek_soru_satirlari("oncelikli_tablo",
                               [("Unexplained fever or weight loss with extrapulmonary findings",
                                 "Açıklanamayan ateş/kilo kaybı + lenfadenopati/plevral efüzyon/kemik-eklem vb. ekstra-pulmoner bulgular"),
                                ("Prolonged cough with constitutional symptoms", "Uzamış öksürük + konstitüsyonel semptom"),
                                ("Both", "Her ikisi"),
                                ("Don't know", "Bilmiyorum")])
satirlar.append(("bolum", "Preferred training format (multiple answers allowed)"))
satirlar += coklu_soru_satirlari([("Written algorithm or protocol", "G9_algoritma"),
                                  ("Case-based face-to-face meeting", "G9_vaka_temelli"),
                                  ("Short online training (20–30 min)", "G9_cevrimici")])
satirlar.append(("bolum", "Topic most in need of a standard algorithm"))
satirlar += tek_soru_satirlari("algoritma_ihtiyaci",
                               [("Whom to screen", "Kime tarama yapılacağı"),
                                ("Isolation and patient flow in suspected active TB", "Aktif TB şüphesinde izolasyon/akış"),
                                ("Choice and monitoring of LTBI treatment", "LTBI tedavisi seçimi ve izlem"),
                                ("Interpretation of IGRA/TST", "IGRA/TST yorumlama"),
                                ("Contact management", "Temaslı yönetimi")])

tablo_olustur(
    belge, 4,
    "Reported LTBI screening practice, management of suspected active TB and training needs",
    ["Question and response", f"n (%)"],
    satirlar,
    "Percentages are calculated among participants who answered each question. "
    "HD, haemodialysis; IGRA, interferon-gamma release assay; LTBI, latent tuberculosis infection; "
    "TB, tuberculosis; TST, tuberculin skin test.",
    [12.5, 3.0],
)


belge.save(tablo_klasoru / "Tablolar.docx")
print("tablolar kaydedildi:", tablo_klasoru / "Tablolar.docx")
